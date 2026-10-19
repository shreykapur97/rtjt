# Copyright (C) 2026 Shrey Kapur
# SPDX-License-Identifier: AGPL-3.0-or-later

"""Publish faculty review notes from the reviews folder to Google Drive as .docx.

Usage: python -m pipeline.publish [--date YYYY-MM-DD]     (no date = all dates)

    <REVIEWS_DIR>/notes/<date>/faculty/<Name>.md  ->  <FACULTY_NOTES_DIR>/<date>/<Name>.docx

FACULTY_NOTES_DIR is set in .env: the Drive for Desktop folder path (not a .lnk shortcut),
outside dev and outside the reviews folder. Only notes/<date>/faculty/ is read. A .docx is only
rebuilt when its note is newer, and nothing in Drive is ever deleted.
"""
import argparse
import os
import re
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .config import ROOT, load_config
from .export import reviews_dir

INLINE = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`|\[([^\]]+)\]\((https?://[^)\s]+)\)")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        props.append(el)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str) -> None:
    """Bold, `code` and [links](url); everything else as plain text."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        bold, code, label, url = m.groups()
        if bold:
            paragraph.add_run(bold).bold = True
        elif code:
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            add_hyperlink(paragraph, label, url)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            para = table.cell(i, j).paragraphs[0]
            text = row[j] if j < len(row) else ""
            if i == 0:
                para.add_run(text).bold = True
            else:
                add_inline(para, text)


def md_to_docx(md: str, dest: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            para = doc.add_paragraph()
            run = para.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"(#{1,4})\s+(.*)", stripped):
            doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif m := re.match(r"[-*]\s+(.*)", stripped):
            style = "List Bullet 2" if line.startswith(("  ", "\t")) else "List Bullet"
            add_inline(doc.add_paragraph(style=style), m.group(1))
        elif m := re.match(r"\d+\.\s+(.*)", stripped):
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
        elif stripped:
            add_inline(doc.add_paragraph(), stripped)
        i += 1
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def faculty_dir(reviews: Path) -> Path:
    value = os.getenv("FACULTY_NOTES_DIR")
    if not value:
        raise SystemExit("Set FACULTY_NOTES_DIR in .env (the Drive for Desktop folder path)")
    path = Path(value).resolve()
    if not path.is_dir():
        raise SystemExit(f"Faculty folder not found (is Google Drive running?): {path}")
    for inside in (ROOT, reviews):
        if path == inside or inside in path.parents:
            raise SystemExit(f"FACULTY_NOTES_DIR must be outside {inside}.")
    return path


def publish(notes: Path, target: Path, date: str | None) -> dict:
    counts = {"new": 0, "updated": 0, "unchanged": 0}
    for src in sorted(notes.glob(f"{date or '*'}/faculty/*.md")):
        day = src.parent.parent.name
        dest = target / day / f"{src.stem}.docx"
        if dest.exists() and dest.stat().st_mtime >= src.stat().st_mtime:
            counts["unchanged"] += 1
            continue
        status = "updated" if dest.exists() else "new"
        md_to_docx(src.read_text(encoding="utf-8"), dest)
        counts[status] += 1
        print(f"{status:8s} {day}/{src.stem}.docx")
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description="Publish faculty review notes to Google Drive as .docx")
    parser.add_argument("--date", help="one review date (YYYY-MM-DD); default all dates")
    args = parser.parse_args()

    load_config()  # loads .env
    reviews = reviews_dir()
    target = faculty_dir(reviews)
    c = publish(reviews / "notes", target, args.date)
    print(f"faculty: {c['new']} new, {c['updated']} updated, {c['unchanged']} unchanged -> {target}")


if __name__ == "__main__":
    main()
